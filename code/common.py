# -*- coding: utf-8 -*-
"""2026 CUMCM C题 微网-外网购电优化 —— 共享模型核心 (common.py)

模型: 线性规划 (LP, scipy linprog / HiGHS)。
目标: 最小化购电费, 满足 供电≥负载、储能容量/功率边界、储能跨日连续。

统一约定 (与 planning/symbol_table.md 一致):
  * 一天 T=144 个 10 分钟区间, Δt=1/6 h。
  * x_t 计划购电量(kWh), c_t 充电量(kWh), q_t 放电量(kWh), E_t 储能量(kWh)。
  * 区间电量 = 区间起点功率 × Δt (A5/HD1)。
  * 充/放电效率单向 η_c=η_d=0.9 (A6/HD2)。
  * 能量平衡: x_t + G_t·Δt + q_t ≥ L_t·Δt + c_t  (A7)。
  * 负荷 L 视为已知(附件2实际), 光伏 G 有预报(附件3)/实际(附件2)。
"""
import os
import pickle
from datetime import datetime

import numpy as np
from scipy.optimize import linprog

SEED = 2026
np.random.seed(SEED)

# ---- 物理/费用参数 (data_clean/cleaned_data.pkl 的 params) ----
T = 144
DT = 1.0 / 6.0
EMIN = 1200.0
EMAX = 10800.0
PMAX = 5000.0
PMAX_E = PMAX * DT            # 单区间最大充/放电量 kWh
ETA_C = 0.9
ETA_D = 0.9
E0_INIT = 6000.0
ALPHA_EM = 5.0                # 紧急购电 5 倍
BETA_PEN = 0.5                # 违约 50%
GAMMA_PRE = 1.5               # 溢价 1.5 倍

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_CLEAN = os.path.join(ROOT, "data_clean")
TEMPLATE_DIR = os.path.join(ROOT, "data_raw", "附件5")
RESULTS_DIR = os.path.join(ROOT, "results")

# 报告期: 2025-02-01 ~ 2025-12-31, 在 365 天中的索引 31..364 (1 月为预热, HD5)
WARMUP_DAYS = 31
REPORT_DAYS = 334


def load_bundle():
    with open(os.path.join(DATA_CLEAN, "cleaned_data.pkl"), "rb") as f:
        return pickle.load(f)


def report_dates(bundle):
    """报告期 334 天对应的 datetime 列表 (与结果模板 日期 列一致)。"""
    ds = bundle["dates"]  # ['2025-01-01 00:00:00', ...]
    out = []
    for s in ds[WARMUP_DAYS:]:
        out.append(datetime.strptime(s, "%Y-%m-%d %H:%M:%S"))
    return out


def time_labels():
    """144 个区间标签, 与结果模板 计划购电量 列头一致。"""
    labels = []
    for i in range(T):
        start_min = (i + 1) * 10
        if start_min == 1440:  # 24:00 -> 0:00+1
            labels.append("0:00+1-0:10+1")
            continue
        sh, sm = divmod(start_min, 60)
        eh, em = divmod(start_min + 10, 60)
        labels.append(f"{sh}:{sm:02d}-{eh}:{em:02d}")
    return labels


SEG_LABELS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
              "12:00-16:00", "16:00-20:00", "20:00-24:00"]


def clip0(a, tol=1e-8):
    return np.where(np.abs(a) < tol, 0.0, np.maximum(a, 0.0))


# ======================================================================
# 1) 单日 LP —— Q1 及逐日滚动的基础
# ======================================================================
def daily_lp(price, load, pv, E0=None, terminal_eq=False, E_term=None):
    """单日线性规划。

    返回 (cost, x, c, q, E), 其中 x/c/q 长度 T, E 长度 T+1。
    terminal_eq=True 时固定 E_T=E_term(缺省 E0)。
    """
    n = 3 * T + (T + 1)

    def ix(t): return t
    def ic(t): return T + t
    def iq(t): return 2 * T + t
    def iE(t): return 3 * T + t

    cobj = np.zeros(n)
    cobj[0:T] = price

    A_ub, b_ub = [], []
    for t in range(T):
        row = np.zeros(n)
        row[ix(t)] = -1.0
        row[iq(t)] = -1.0
        row[ic(t)] = 1.0
        A_ub.append(row); b_ub.append(pv[t] * DT - load[t] * DT)
    for t in range(T):
        row = np.zeros(n); row[ic(t)] = 1.0
        A_ub.append(row); b_ub.append(PMAX_E)
    for t in range(T):
        row = np.zeros(n); row[iq(t)] = 1.0
        A_ub.append(row); b_ub.append(PMAX_E)

    A_eq, b_eq = [], []
    for t in range(T):
        row = np.zeros(n)
        row[iE(t + 1)] = 1.0
        row[iE(t)] = -1.0
        row[ic(t)] = -ETA_C
        row[iq(t)] = 1.0 / ETA_D
        A_eq.append(row); b_eq.append(0.0)
    row = np.zeros(n); row[iE(0)] = 1.0
    A_eq.append(row); b_eq.append(E0 if E0 is not None else E0_INIT)
    if terminal_eq:
        row = np.zeros(n); row[iE(T)] = 1.0
        A_eq.append(row); b_eq.append(E_term if E_term is not None else (E0 if E0 is not None else E0_INIT))

    bounds = [(0, None)] * (3 * T) + [(EMIN, EMAX)] * (T + 1)
    res = linprog(cobj, A_ub=np.array(A_ub), b_ub=np.array(b_ub),
                  A_eq=np.array(A_eq), b_eq=np.array(b_eq),
                  bounds=bounds, method="highs")
    if not res.success:
        raise RuntimeError("daily_lp infeasible: " + res.message)
    v = res.x
    x = clip0(v[0:T]); c = clip0(v[T:2 * T]); q = clip0(v[2 * T:3 * T])
    E = v[3 * T:3 * T + T + 1]
    return float(res.fun), x, c, q, E


# ======================================================================
# 2) 全年联合 LP —— Q2/Q4 的精确解 (消除逐日滚动末端短视)
# ======================================================================
def full_year_lp(price_flat, load_flat, pv_flat, E0=None, terminal_eq=False, E_term=None):
    """全年联合线性规划 (稀疏, HiGHS)。

    变量: x[N], c[N], q[N], E[N+1];  N = len(price_flat) = D*T。
    返回 (cost, X(D,T), C(D,T), Q(D,T), E(N+1,))。
    """
    E0 = E0_INIT if E0 is None else E0
    N = len(price_flat)
    D = N // T
    n = 4 * N + 1

    def ix(k): return k
    def ic(k): return N + k
    def iq(k): return 2 * N + k
    def iE(k): return 3 * N + k

    cobj = np.zeros(n)
    cobj[:N] = price_flat

    rows, cols, data = [], [], []
    b_ub = np.empty(3 * N)
    # 能量平衡 -x-q+c <= pv*dt - load*dt
    for k in range(N):
        r = k
        rows += [r, r, r]; cols += [ix(k), iq(k), ic(k)]; data += [-1.0, -1.0, 1.0]
        b_ub[r] = pv_flat[k] * DT - load_flat[k] * DT
    # 充电功率 c <= PmaxE
    for k in range(N):
        r = N + k
        rows.append(r); cols.append(ic(k)); data.append(1.0); b_ub[r] = PMAX_E
    # 放电功率 q <= PmaxE
    for k in range(N):
        r = 2 * N + k
        rows.append(r); cols.append(iq(k)); data.append(1.0); b_ub[r] = PMAX_E

    from scipy import sparse
    A_ub = sparse.coo_matrix((data, (rows, cols)), shape=(3 * N, n)).tocsr()

    rows2, cols2, data2 = [], [], []
    b_eq_list = [0.0] * N
    # 储能动态 E_{k+1}-E_k - ηc c_k + q_k/ηd = 0
    for k in range(N):
        r = k
        rows2 += [r, r, r, r]
        cols2 += [iE(k + 1), iE(k), ic(k), iq(k)]
        data2 += [1.0, -1.0, -ETA_C, 1.0 / ETA_D]
    r = N
    rows2.append(r); cols2.append(iE(0)); data2.append(1.0)
    b_eq_list.append(E0)
    if terminal_eq:
        r = N + 1
        rows2.append(r); cols2.append(iE(N)); data2.append(1.0)
        b_eq_list.append(E_term if E_term is not None else E0)
    b_eq = np.array(b_eq_list)
    n_eq = N + 1 + (1 if terminal_eq else 0)
    A_eq = sparse.coo_matrix((data2, (rows2, cols2)), shape=(n_eq, n)).tocsr()

    bounds = [(0, None)] * (3 * N) + [(EMIN, EMAX)] * (N + 1)
    res = linprog(cobj, A_ub=A_ub, b_ub=b_ub, A_eq=A_eq, b_eq=b_eq,
                  bounds=bounds, method="highs")
    if not res.success:
        raise RuntimeError("full_year_lp infeasible: " + res.message)
    v = res.x
    X = clip0(v[:N]).reshape(D, T)
    C = clip0(v[N:2 * N]).reshape(D, T)
    Q = clip0(v[2 * N:3 * N]).reshape(D, T)
    E = v[3 * N:]
    return float(res.fun), X, C, Q, E


# ======================================================================
# 3) 基线: 无储能 / 规则套利
# ======================================================================
def no_storage(price, load, pv):
    """无储能基线: x_t = max(L·Δt − G·Δt, 0)。返回 (cost, x)。"""
    x = np.maximum(load * DT - pv * DT, 0.0)
    return float((price * x).sum()), x


def no_storage_dayahead(price, load, pv_actual, forecast, d):
    """无储能 + 日前预报 + 紧急购电基线 (与 M2 同信息, 公平对比)。

    0:00 计划 x = max(L·Δt − ghat0·Δt, 0) (按预报买缺口);
    紧急 z = max(min(ghat0, L) − G^实际, 0)·Δt (预报高估光伏时欠购, 5倍)。
    返回 (cost, x, z)。
    """
    ghat0 = forecast_intervals(forecast, d, 0)
    x = np.maximum(load * DT - ghat0 * DT, 0.0)
    z = np.maximum(np.minimum(ghat0, load) * DT - pv_actual * DT, 0.0)
    cost = float((price * x).sum() + ALPHA_EM * (price * z).sum())
    return cost, x, z


def rule_arbitrage(price, load, pv, E0=EMIN, terminal_eq=False, E_term=None):
    """规则峰谷套利基线 (贪心时序模拟)。

    低价 1/3 区间充到满、高价 1/3 区间放到空、中价不动; 余下购电补缺口。
    terminal_eq=True 时在最后 NB=min(24,M) 区间按功率/容量上限把储能拉回 E_term(缺省 E0),
    实现与 LP 方案一致的日闭合 E(T)=E_term。
    返回 (cost, x, c, q, E)。
    """
    plo, phi = float(np.min(price)), float(np.max(price))
    th_lo = plo + (phi - plo) / 3.0
    th_hi = plo + 2.0 * (phi - plo) / 3.0
    E = float(E0)
    M = len(price)
    x = np.zeros(M); c = np.zeros(M); q = np.zeros(M)
    Est = np.zeros(M + 1); Est[0] = E
    for t in range(M):
        if price[t] <= th_lo:
            c[t] = min(PMAX_E, (EMAX - E) / ETA_C)
            q[t] = 0.0
        elif price[t] >= th_hi:
            q[t] = min(PMAX_E, (E - EMIN) * ETA_D)
            c[t] = 0.0
        else:
            c[t] = q[t] = 0.0
        E = E + ETA_C * c[t] - q[t] / ETA_D
        E = min(max(E, EMIN), EMAX)
        Est[t + 1] = E
        x[t] = max(load[t] * DT + c[t] - pv[t] * DT - q[t], 0.0)
    if terminal_eq:
        E_term_val = E_term if E_term is not None else E0
        NB = min(24, M)
        E_cur = Est[M - NB]
        dE = E_term_val - E_cur
        if abs(dE) > 1e-8:
            per = dE / NB
            for t in range(M - NB, M):
                if per > 0.0:            # 需充电拉回
                    c[t] = min(PMAX_E, per / ETA_C)
                    q[t] = 0.0
                else:                     # 需放电拉回
                    c[t] = 0.0
                    q[t] = min(PMAX_E, (-per) * ETA_D)
                E_cur = E_cur + ETA_C * c[t] - q[t] / ETA_D
                E_cur = min(max(E_cur, EMIN), EMAX)
                Est[t + 1] = E_cur
                x[t] = max(load[t] * DT + c[t] - pv[t] * DT - q[t], 0.0)
    cost = float((price * x).sum())
    return cost, x, c, q, Est


# ======================================================================
# 4) Q3 滚动 MPC 组件
# ======================================================================
def forecast_intervals(forecast, d, s_idx):
    """附件3 预报 -> 144 区间光伏功率(kW)。

    issue s_idx∈{0,1,2,3} 对应 0/6/12/18 点; 覆盖小时 [6*s_idx, 24)。
    区间 t 属于小时 h=t//6, 取值 forecast[d*4+s_idx, h-6*s_idx]。
    """
    ghat = np.zeros(T)
    issue_hour = 6 * s_idx
    for h in range(issue_hour, 24):
        val = forecast[d * 4 + s_idx, h - issue_hour]
        ghat[6 * h:6 * h + 6] = val
    return ghat


def pv_overestimation(forecast, pv):
    """预计算各 (天, issue) 的光伏高估 o = Ĝ − G (10min 区间, 每块 36 区间)。

    仅取 MPC 实际执行块: issue s 执行 [36s, 36s+36), 预报取 fc[d*4+s, j//6]。
    返回 (D, 4, 36)。
    """
    D = pv.shape[0]
    over = np.zeros((D, 4, 36))
    for s in range(4):
        fc_s = forecast[s::4]              # (D, 24) issue s
        for j in range(36):
            over[:, s, j] = fc_s[:, j // 6] - pv[:, 36 * s + j]
    return over


def issue_risk_delta(over, d, quantile=80.0, min_hist=7):
    """day d 的各 issue 保守裕量 δ[s] = P_q(光伏高估), 因果(仅用 [min_hist, d) 历史)。

    返回长度 4 的 numpy 数组; d <= min_hist 时全 0。
    """
    delta = np.zeros(4)
    if d <= min_hist:
        return delta
    hist = over[min_hist:d]                # (d-min_hist, 4, 36)
    for s in range(4):
        delta[s] = float(np.percentile(hist[:, s, :], quantile))
    return delta


def stage_lp(price, load, pv, x_plan, E0, terminal_eq=False, E_term=None):
    """Q3 调整阶段 LP (在给定 0:00 计划 x_plan 下, 求剩余时段最优调整 y)。

    分段计费线性化: y = u + v, 0≤u≤x_plan(计划内), v≥0(溢价);
    费用 = p·[0.5x + 0.5u + 1.5v], 常数项 0.5x 略去 -> min Σ p(0.5u+1.5v)。
    terminal_eq=True 时固定 E_M=E_term(缺省 E0), 实现日闭合 E(144)=6000。
    返回 (cost_var, y, c, q, E); y/c/q 长度 M, E 长度 M+1。
    """
    M = len(price)
    n = 4 * M + (M + 1)

    def iu(t): return t
    def iv(t): return M + t
    def ic(t): return 2 * M + t
    def iq(t): return 3 * M + t
    def iE(t): return 4 * M + t

    cobj = np.zeros(n)
    for t in range(M):
        cobj[iu(t)] = 0.5 * price[t]
        cobj[iv(t)] = 1.5 * price[t]

    A_ub, b_ub = [], []
    for t in range(M):
        # 能量平衡: u+v+q-c >= load*dt - pv*dt  => -u-v-q+c <= pv*dt-load*dt
        row = np.zeros(n)
        row[iu(t)] = -1.0; row[iv(t)] = -1.0; row[iq(t)] = -1.0; row[ic(t)] = 1.0
        A_ub.append(row); b_ub.append(pv[t] * DT - load[t] * DT)
    for t in range(M):
        row = np.zeros(n); row[iu(t)] = 1.0   # u ≤ x_plan
        A_ub.append(row); b_ub.append(x_plan[t])
    for t in range(M):
        row = np.zeros(n); row[ic(t)] = 1.0
        A_ub.append(row); b_ub.append(PMAX_E)
    for t in range(M):
        row = np.zeros(n); row[iq(t)] = 1.0
        A_ub.append(row); b_ub.append(PMAX_E)

    A_eq, b_eq = [], []
    for t in range(M):
        row = np.zeros(n)
        row[iE(t + 1)] = 1.0; row[iE(t)] = -1.0
        row[ic(t)] = -ETA_C; row[iq(t)] = 1.0 / ETA_D
        A_eq.append(row); b_eq.append(0.0)
    row = np.zeros(n); row[iE(0)] = 1.0
    A_eq.append(row); b_eq.append(E0)
    if terminal_eq:
        row = np.zeros(n); row[iE(M)] = 1.0
        A_eq.append(row); b_eq.append(E_term if E_term is not None else E0)

    bounds = [(0, None)] * (4 * M) + [(EMIN, EMAX)] * (M + 1)
    res = linprog(cobj, A_ub=np.array(A_ub), b_ub=np.array(b_ub),
                  A_eq=np.array(A_eq), b_eq=np.array(b_eq),
                  bounds=bounds, method="highs")
    if not res.success:
        raise RuntimeError("stage_lp infeasible: " + res.message)
    v = res.x
    u = clip0(v[:M]); vv = clip0(v[M:2 * M])
    c = clip0(v[2 * M:3 * M]); q = clip0(v[3 * M:4 * M])
    E = v[4 * M:]
    y = u + vv
    return float(res.fun), y, c, q, E


def q3_mpc_day(price, load_d, pv_actual, forecast, d, E_carry, delta=None):
    """Q3 单日滚动 MPC。delta: 各 issue 保守裕量(长度4, 单位 kW), None=不修正。

    返回 (cost, x, y_final, c_final, q_final, E_end, z)。
    """
    delta = np.zeros(4) if delta is None else np.asarray(delta, dtype=float)
    # 阶段0 (0:00): 计划 x, 用 0:00 保守预报 (日闭合 E(144)=E(0)=6000)
    ghat0 = clip0(forecast_intervals(forecast, d, 0) - delta[0])
    _, x, c0, q0, E0t = daily_lp(price, load_d, ghat0, E0=E_carry,
                                 terminal_eq=True, E_term=E0_INIT)

    # 阶段1 (6:00): 调整 t∈[36,144)
    ghat1 = clip0(forecast_intervals(forecast, d, 1) - delta[1])
    sl = 36
    _, y1, c1, q1, E1t = stage_lp(price[sl:], load_d[sl:], ghat1[sl:], x[sl:],
                                  E0=E0t[sl], terminal_eq=True, E_term=E0_INIT)

    # 阶段2 (12:00): 调整 t∈[72,144)
    ghat2 = clip0(forecast_intervals(forecast, d, 2) - delta[2])
    sl = 72
    _, y2, c2, q2, E2t = stage_lp(price[sl:], load_d[sl:], ghat2[sl:], x[sl:],
                                  E0=E1t[sl - 36], terminal_eq=True, E_term=E0_INIT)

    # 阶段3 (18:00): 调整 t∈[108,144)
    ghat3 = clip0(forecast_intervals(forecast, d, 3) - delta[3])
    sl = 108
    _, y3, c3, q3, E3t = stage_lp(price[sl:], load_d[sl:], ghat3[sl:], x[sl:],
                                  E0=E2t[sl - 72], terminal_eq=True, E_term=E0_INIT)

    # 拼接最终执行方案
    y_final = np.concatenate([x[:36], y1[:36], y2[:36], y3[:36]])
    c_final = np.concatenate([c0[:36], c1[:36], c2[:36], c3[:36]])
    q_final = np.concatenate([q0[:36], q1[:36], q2[:36], q3[:36]])
    E_end = E3t[-1]

    # 结算: 紧急购电(用附件2实际光伏) + 分段计费
    z = np.maximum(load_d * DT + c_final - y_final - pv_actual * DT - q_final, 0.0)
    z = clip0(z)
    bill = price * (np.minimum(x, y_final)
                    + BETA_PEN * np.maximum(x - y_final, 0.0)
                    + GAMMA_PRE * np.maximum(y_final - x, 0.0))
    cost = float(bill.sum() + ALPHA_EM * (price * z).sum())
    return cost, x, y_final, c_final, q_final, E_end, z


def rule_mpc_day(price, load_d, pv_actual, forecast, d, E_carry, delta=None):
    """Q3 基线 R3 (规则套利 + 滚动预报): 与 M3 同信息(滚动预报)、同计费, 仅用贪心规则替代 LP。

    delta: 各 issue 保守裕量(长度4), None=不修正。
    返回 (cost, x, y_final, c_final, q_final, E_end, z)。"""
    delta = np.zeros(4) if delta is None else np.asarray(delta, dtype=float)
    ghat0 = clip0(forecast_intervals(forecast, d, 0) - delta[0])
    _, x, c0, q0, E0t = rule_arbitrage(price, load_d, ghat0, E0=E_carry,
                                       terminal_eq=True, E_term=E0_INIT)

    ghat1 = clip0(forecast_intervals(forecast, d, 1) - delta[1])
    sl = 36
    _, y1, c1, q1, E1t = rule_arbitrage(price[sl:], load_d[sl:], ghat1[sl:], E0=E0t[sl],
                                        terminal_eq=True, E_term=E0_INIT)

    ghat2 = clip0(forecast_intervals(forecast, d, 2) - delta[2])
    sl = 72
    _, y2, c2, q2, E2t = rule_arbitrage(price[sl:], load_d[sl:], ghat2[sl:], E0=E1t[sl - 36],
                                        terminal_eq=True, E_term=E0_INIT)

    ghat3 = clip0(forecast_intervals(forecast, d, 3) - delta[3])
    sl = 108
    _, y3, c3, q3, E3t = rule_arbitrage(price[sl:], load_d[sl:], ghat3[sl:], E0=E2t[sl - 72],
                                        terminal_eq=True, E_term=E0_INIT)

    y_final = np.concatenate([x[:36], y1[:36], y2[:36], y3[:36]])
    c_final = np.concatenate([c0[:36], c1[:36], c2[:36], c3[:36]])
    q_final = np.concatenate([q0[:36], q1[:36], q2[:36], q3[:36]])
    E_end = E3t[-1]

    z = np.maximum(load_d * DT + c_final - y_final - pv_actual * DT - q_final, 0.0)
    z = clip0(z)
    bill = price * (np.minimum(x, y_final)
                    + BETA_PEN * np.maximum(x - y_final, 0.0)
                    + GAMMA_PRE * np.maximum(y_final - x, 0.0))
    cost = float(bill.sum() + ALPHA_EM * (price * z).sum())
    return cost, x, y_final, c_final, q_final, E_end, z


def q3_static_day(price, load_d, pv_actual, forecast, d, E_carry, delta=None):
    """Q3 基线 B3 (静态不调整): 仅 0:00 计划并执行, 预报误差全由紧急购电吸收。

    delta: 各 issue 保守裕量(长度4), None=不修正。
    返回 (cost, x, c, q, E_end, z)。"""
    delta = np.zeros(4) if delta is None else np.asarray(delta, dtype=float)
    ghat0 = clip0(forecast_intervals(forecast, d, 0) - delta[0])
    _, x, c, q, Et = daily_lp(price, load_d, ghat0, E0=E_carry,
                              terminal_eq=True, E_term=E0_INIT)
    z = np.maximum(load_d * DT + c - x - pv_actual * DT - q, 0.0)
    z = clip0(z)
    cost = float((price * x).sum() + ALPHA_EM * (price * z).sum())
    return cost, x, c, q, Et[-1], z


# ======================================================================
# 5) 结果写盘 (加载模板, 填充数据; 骨架 sheet 重建为完整 334 天)
# ======================================================================
def fill_plan_grid(ws, X):
    """填充 计划购电量/调整购电量 网格 (335×144)。"""
    for d in range(REPORT_DAYS):
        for t in range(T):
            ws.cell(row=2 + d, column=2 + t, value=float(X[d, t]))


def rebuild_charge_sheet(ws, dates, C, Q, E):
    """重建 充放电量 sheet: 334 天 × 6 段, 0:00/24:00 储电量。"""
    import openpyxl
    ws.delete_rows(1, ws.max_row)
    hdr = ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=1, column=j, value=h)
    r = 2
    for d in range(REPORT_DAYS):
        E0_day = float(E[d * T])
        E24 = float(E[d * T + T])
        for s in range(6):
            lo, hi = s * 24, (s + 1) * 24
            chg = float(C[d, lo:hi].sum())
            dis = float(Q[d, lo:hi].sum())
            ws.cell(row=r, column=1, value=dates[d] if s == 0 else None)
            ws.cell(row=r, column=2, value=SEG_LABELS[s])
            ws.cell(row=r, column=3, value=chg)
            ws.cell(row=r, column=4, value=dis)
            if s == 0:
                ws.cell(row=r, column=5, value=datetime(1900, 1, 1, 0, 0).time())
                ws.cell(row=r, column=6, value=E0_day)
            elif s == 1:
                ws.cell(row=r, column=5, value="24:00")
                ws.cell(row=r, column=6, value=E24)
            r += 1


def rebuild_emergency_sheet(ws, dates, Z, labels):
    """重建 紧急购电量 sheet (稀疏: 仅 z>0 的区间)。"""
    ws.delete_rows(1, ws.max_row)
    hdr = ["日期", "购电时间段", "购电量"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=1, column=j, value=h)
    r = 2
    for d in range(REPORT_DAYS):
        for t in range(T):
            if Z[d, t] > 1e-8:
                ws.cell(row=r, column=1, value=dates[d])
                ws.cell(row=r, column=2, value=labels[t])
                ws.cell(row=r, column=3, value=float(Z[d, t]))
                r += 1
    # 若全年无紧急购电(完美预见), 保留一行占位说明
    if r == 2:
        ws.cell(row=2, column=1, value=dates[0])
        ws.cell(row=2, column=2, value="无")
        ws.cell(row=2, column=3, value=0.0)
