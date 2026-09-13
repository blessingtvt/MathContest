# -*- coding: utf-8 -*-
# Q4 主模型 M4 (round3, 块级δ): 波动电价下完全继承 Q2/Q3 最终模型, 电价 p_t → p_{d,t}。
# result4-2 (对应 Q2): 逐日滚动 LP + 附件2净负荷自预报(7天加权) + 按星期分桶80分位修正 + 日闭合。
# result4-3 (对应 Q3): 滚动 MPC + 分段计费 + 日闭合 + issue×执行块 逐区间80分位光伏修正 δ(4,36)。
# 基线: B4 无储能 / R4 规则套利 / B4_ref 无储能完美预见(诊断下界)。
import os, sys, json
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import numpy as np
import openpyxl
from common import *

ROUND = "round3"
QUANTILE = 80.0
MIN_HIST = 7
W = [0.3, 0.2, 0.15, 0.12, 0.1, 0.08, 0.05]


def bundle_dow(bundle):
    """365 天的星期索引 (0=周一..6=周日)，与 bundle['dates'] 对齐。"""
    return np.array([datetime.strptime(s, "%Y-%m-%d %H:%M:%S").weekday()
                     for s in bundle["dates"]])


def error_quantile_dow(Eerr, full_dow, d, quantile=80.0):
    """按星期分桶的历史误差 quantile 分位 (仅用 d 之前同星期数据, 因果)。"""
    hist = Eerr[7:d]
    hd = full_dow[7:d]
    mask = hd == full_dow[d]
    if mask.sum() == 0:
        return np.zeros(Eerr.shape[1])
    return np.percentile(hist[mask], quantile, axis=0)


def netload_forecast(N):
    """7 天加权平均净负荷预报 (仅用历史, 因果)。返回 (Nhat, Eerr)。"""
    D = N.shape[0]
    w = np.array(W)
    Nhat = np.zeros_like(N)
    for d in range(len(w), D):
        for k in range(len(w)):
            Nhat[d] += w[k] * N[d - 1 - k]
    return Nhat, N - Nhat


def no_storage_forecast(price, Nsafe, Nreal):
    """B4 基线: 无储能 + 自预报(含风险修正), 缺口紧急购电。返回 (cost, x, z)。"""
    x = np.maximum(Nsafe * DT, 0.0)
    z = np.maximum(Nreal * DT - x, 0.0)
    cost = float((price * x).sum() + ALPHA_EM * (price * z).sum())
    return cost, x, z


def rule_forecast_day(price, Nsafe, Nreal, E_carry):
    """R4 基线: 规则套利(价格阈值) + 自预报, 缺口紧急购电。返回 (cost, x, c, q, E_end, z)。"""
    _, x, c, q, Est = rule_arbitrage(price, Nsafe, np.zeros(T), E0=E_carry)
    z = clip0(np.maximum(Nreal * DT + c - x - q, 0.0))
    cost = float((price * x).sum() + ALPHA_EM * (price * z).sum())
    return cost, x, c, q, Est[-1], z


def block_delta(over, d, quantile=80.0, min_hist=7):
    """issue×执行块 逐区间 分位修正 (4,36)，因果(仅用 [min_hist, d) 历史)。"""
    hist = over[min_hist:d]
    if len(hist) == 0:
        return np.zeros((4, 36))
    qv = np.asarray(quantile if np.ndim(quantile) else [quantile] * 4, dtype=float)
    return np.stack([np.percentile(hist[:, s, :], qv[s], axis=0) for s in range(4)])


def corrected_forecast_block(fc, d, s, delta):
    """块级修正：只对 issue s 实际执行的 6h 块 [36s, 36s+36) 减去 delta[s](36 向量)。"""
    g = forecast_intervals(fc, d, s).copy()
    start = 36 * s
    g[start:start + 36] = clip0(g[start:start + 36] - delta[s])
    return g


def q3_mpc_day_block(price, load_d, pv_actual, fc, d, E_carry, delta):
    """Q3 单日滚动 MPC（块级 δ，delta 形状 (4,36)）。返回 (cost, x, y, c, q, E_end, z, E_day)。"""
    g0 = corrected_forecast_block(fc, d, 0, delta)
    _, x, c0, q0, E0t = daily_lp(price, load_d, g0, E0=E_carry, terminal_eq=True, E_term=E0_INIT)
    g1 = corrected_forecast_block(fc, d, 1, delta); sl = 36
    _, y1, c1, q1, E1t = stage_lp(price[sl:], load_d[sl:], g1[sl:], x[sl:], E0t[sl], terminal_eq=True, E_term=E0_INIT)
    g2 = corrected_forecast_block(fc, d, 2, delta); sl = 72
    _, y2, c2, q2, E2t = stage_lp(price[sl:], load_d[sl:], g2[sl:], x[sl:], E1t[sl - 36], terminal_eq=True, E_term=E0_INIT)
    g3 = corrected_forecast_block(fc, d, 3, delta); sl = 108
    _, y3, c3, q3, E3t = stage_lp(price[sl:], load_d[sl:], g3[sl:], x[sl:], E2t[sl - 72], terminal_eq=True, E_term=E0_INIT)
    y = np.concatenate([x[:36], y1[:36], y2[:36], y3[:36]])
    c = np.concatenate([c0[:36], c1[:36], c2[:36], c3[:36]])
    qv = np.concatenate([q0[:36], q1[:36], q2[:36], q3[:36]])
    z = clip0(load_d * DT + c - y - pv_actual * DT - qv)
    bill = price * (np.minimum(x, y) + BETA_PEN * np.maximum(x - y, 0.0) + GAMMA_PRE * np.maximum(y - x, 0.0))
    E_day = np.concatenate([E0t[:37], E1t[1:37], E2t[1:37], E3t[1:37]])
    return float(bill.sum() + ALPHA_EM * (price * z).sum()), x, y, c, qv, E3t[-1], z, E_day


def fill_daily_summary(ws, Qgrid, daily_cost):
    """填充 全天购电量(第146列, kWh) 与 全天购电费(第147列, 元)。"""
    for d in range(REPORT_DAYS):
        ws.cell(row=2 + d, column=146, value=float(Qgrid[d].sum()))
        ws.cell(row=2 + d, column=147, value=float(daily_cost[d]))


def main():
    b = load_bundle()
    load = b["附件2_load"]
    pv = b["附件2_pv"]
    fc = b["附件3"]["forecast"]
    price_d = b["附件4_price"]          # (365,144) 波动电价
    D = load.shape[0]
    dates = report_dates(b)
    labels = time_labels()
    N = load - pv                        # 净负荷 (kW)
    Nhat, Eerr = netload_forecast(N)
    full_dow = bundle_dow(b)
    over = pv_overestimation(fc, pv)

    R = REPORT_DAYS
    out_dir = os.path.join(RESULTS_DIR, "Q4")
    os.makedirs(os.path.join(out_dir, "experiments", ROUND), exist_ok=True)

    # result4-2: Q2 模型 + 波动电价
    Xr = np.zeros((R, T)); Cr = np.zeros((R, T)); Qr = np.zeros((R, T)); Zr = np.zeros((R, T))
    E_seq = [E0_INIT]
    for i, d in enumerate(range(WARMUP_DAYS, D)):
        e80 = error_quantile_dow(Eerr, full_dow, d, QUANTILE)
        Nsafe = Nhat[d] + e80
        _, x_d, c_d, q_d, E_d = daily_lp(price_d[d], Nsafe, np.zeros(T), E0=E0_INIT,
                                         terminal_eq=True, E_term=E0_INIT)
        z_d = clip0(np.maximum(N[d] * DT + c_d - x_d - q_d, 0.0))
        Xr[i] = x_d; Cr[i] = c_d; Qr[i] = q_d; Zr[i] = z_d
        E_seq.extend(E_d[1:].tolist())
    Er = np.array(E_seq)                 # 长度 R*T+1
    plan42 = float((price_d[WARMUP_DAYS:] * Xr).sum())
    emg42 = float(ALPHA_EM * (price_d[WARMUP_DAYS:] * Zr).sum())
    cost42 = plan42 + emg42

    cost42_daily = (price_d[WARMUP_DAYS:] * Xr).sum(axis=1) \
        + ALPHA_EM * (price_d[WARMUP_DAYS:] * Zr).sum(axis=1)

    wb = openpyxl.load_workbook(os.path.join(TEMPLATE_DIR, "result4-2.xlsx"))
    fill_plan_grid(wb["计划购电量"], Xr)
    fill_daily_summary(wb["计划购电量"], Xr, cost42_daily)
    rebuild_charge_sheet(wb["充放电量"], dates, Cr, Qr, Er)
    rebuild_emergency_sheet(wb["紧急购电量"], dates, Zr, labels)
    wb.save(os.path.join(out_dir, "result4-2.xlsx"))

    # result4-3: Q3 模型(块级 δ) + 波动电价
    X3 = np.zeros((R, T)); Y3 = np.zeros((R, T)); C3 = np.zeros((R, T))
    Q3 = np.zeros((R, T)); Z3 = np.zeros((R, T))
    E0_arr = np.zeros(R); E24_arr = np.zeros(R)
    E3_seq = [E0_INIT]
    E = E0_INIT
    row = 0
    for d in range(D):
        delta = block_delta(over, d, QUANTILE, MIN_HIST)
        E_start = E
        cost_d, x, yf, cf, qf, E_end, z, E_day = q3_mpc_day_block(price_d[d], load[d], pv[d], fc, d, E, delta)
        E = E_end
        if d >= WARMUP_DAYS:
            E0_arr[row] = E_start; E24_arr[row] = E_end
            X3[row], Y3[row], C3[row], Q3[row], Z3[row] = x, yf, cf, qf, z
            E3_seq.extend(E_day[1:].tolist())
            row += 1
    E3_flat = np.array(E3_seq)             # 长度 R*T+1, 真实储能轨迹
    plan43 = float((price_d[WARMUP_DAYS:] * np.minimum(X3, Y3)).sum())
    pen43 = float((price_d[WARMUP_DAYS:] * (BETA_PEN * np.maximum(X3 - Y3, 0.0)
                                            + GAMMA_PRE * np.maximum(Y3 - X3, 0.0))).sum())
    emg43 = float(ALPHA_EM * (price_d[WARMUP_DAYS:] * Z3).sum())
    cost43 = plan43 + pen43 + emg43

    cost43_daily = (price_d[WARMUP_DAYS:] * np.minimum(X3, Y3)).sum(axis=1) \
        + (price_d[WARMUP_DAYS:] * (BETA_PEN * np.maximum(X3 - Y3, 0.0)
                                    + GAMMA_PRE * np.maximum(Y3 - X3, 0.0))).sum(axis=1) \
        + ALPHA_EM * (price_d[WARMUP_DAYS:] * Z3).sum(axis=1)

    wb = openpyxl.load_workbook(os.path.join(TEMPLATE_DIR, "result4-3.xlsx"))
    fill_plan_grid(wb["计划购电量"], X3)
    fill_plan_grid(wb["调整购电量"], Y3)
    fill_daily_summary(wb["计划购电量"], X3, cost43_daily)
    fill_daily_summary(wb["调整购电量"], Y3, cost43_daily)
    rebuild_charge_sheet(wb["充放电量"], dates, C3, Q3, E3_flat)
    rebuild_emergency_sheet(wb["紧急购电量"], dates, Z3, labels)
    wb.save(os.path.join(out_dir, "result4-3.xlsx"))

    # 基线 (波动电价, 同信息集)
    b4 = b4ref = r4 = 0.0
    Erule = E0_INIT
    for d in range(WARMUP_DAYS, D):
        e80 = error_quantile_dow(Eerr, full_dow, d, QUANTILE)
        Nsafe = Nhat[d] + e80
        b4 += no_storage_forecast(price_d[d], Nsafe, N[d])[0]
        b4ref += no_storage(price_d[d], load[d], pv[d])[0]
        rc, _, _, _, E_end, _ = rule_forecast_day(price_d[d], Nsafe, N[d], Erule)
        r4 += rc
        Erule = E_end

    # 诊断参照: 储能 + 完美预见 (A 日闭合 / B 自由结转)
    # A: 与 result4-2 同约束(日闭合) 但完美预见真实净负荷, 隔离储能侧预报误差成本
    cost_A = 0.0
    for d in range(WARMUP_DAYS, D):
        cA, _, _, _, _ = daily_lp(price_d[d], N[d], np.zeros(T), E0=E0_INIT,
                                  terminal_eq=True, E_term=E0_INIT)
        cost_A += cA
    # B: 全年自由结转(无日闭合), 全局最优理论下界
    pf = price_d[WARMUP_DAYS:].ravel(); lf = load[WARMUP_DAYS:].ravel(); gv = pv[WARMUP_DAYS:].ravel()
    cost_B, _, _, _, _ = full_year_lp(pf, lf, gv, E0=E0_INIT, terminal_eq=False)

    # 指标
    pv_report = float(pv[WARMUP_DAYS:].sum() * DT)
    emg_energy42 = float(Zr.sum())
    curt42 = float(clip0(Xr + pv[WARMUP_DAYS:] * DT + Qr - (load[WARMUP_DAYS:] * DT + Cr)).sum())
    emg_energy43 = float(Z3.sum())
    curt43 = float(clip0(Y3 + pv[WARMUP_DAYS:] * DT + Q3 - (load[WARMUP_DAYS:] * DT + C3)).sum())

    summary = {
        "schema_version": 1,
        "question_id": "Q4",
        "experiment_id": ROUND,
        "implementation_target": "python",
        "method": "M4-inherit-Q2/Q3-final+fluctuating-price",
        "seed": SEED,
        "approved_decision_id": "q4_method_choice_r2",
        "status": "success",
        "created_at": datetime.now().isoformat(),
        "data_sources": ["data_clean/cleaned_data.pkl"],
        "report_period": "2025-02-01 ~ 2025-12-31",
        "result4_2": {
            "objective_name": "报告期总购电费(元, Q2模型+波动电价)",
            "objective_value": round(cost42, 2),
            "plan_cost": round(plan42, 2),
            "emergency_cost": round(emg42, 2),
            "outputs": ["results/Q4/result4-2.xlsx"],
        },
        "result4_3": {
            "objective_name": "报告期总购电费(元, Q3模型+波动电价)",
            "objective_value": round(cost43, 2),
            "plan_cost": round(plan43, 2),
            "penalty_cost": round(pen43, 2),
            "emergency_cost": round(emg43, 2),
            "outputs": ["results/Q4/result4-3.xlsx"],
        },
        "baselines": [
            {"id": "B4", "name": "无储能(自预报+星期80分位,波动电价)", "objective_value": round(b4, 2)},
            {"id": "R4", "name": "规则套利(自预报+星期80分位,波动电价)", "objective_value": round(r4, 2)},
            {"id": "B4_ref", "name": "无储能完美预见下界(诊断基准,非可执行策略)", "objective_value": round(b4ref, 2)},
        ],
        "diagnostics": [
            {"id": "A", "name": "储能完美预见(日闭合,与result4-2同约束)", "objective_value": round(cost_A, 2)},
            {"id": "B", "name": "储能完美预见(自由结转,全局最优理论下界)", "objective_value": round(cost_B, 2)},
        ],
        "metrics": {
            "result4_2": {
                "储能价值_元": round(b4 - cost42, 2),
                "优化价值_元": round(r4 - cost42, 2),
                "预报误差成本_无储能侧_元": round(b4 - b4ref, 2),
                "预报误差成本_储能侧_元": round(cost42 - cost_A, 2),
                "日闭合约束成本_元": round(cost_A - cost_B, 2),
                "完美预见总价值_储能侧_元": round(cost42 - cost_B, 2),
                "紧急购电_电量kWh": round(emg_energy42, 2),
                "弃光_电量kWh": round(curt42, 2),
                "弃光率": round(curt42 / pv_report, 6),
                "储电量_min": round(float(Er.min()), 2),
                "储电量_max": round(float(Er.max()), 2),
            },
            "result4_3": {
                "紧急购电_电量kWh": round(emg_energy43, 2),
                "弃光_电量kWh": round(curt43, 2),
                "弃光率": round(curt43 / pv_report, 6),
                "储电量_min": round(float(E3_flat.min()), 2),
                "储电量_max": round(float(E3_flat.max()), 2),
                "储电量_E0_min": round(float(E0_arr.min()), 2),
                "储电量_E0_max": round(float(E0_arr.max()), 2),
                "储电量_E24_min": round(float(E24_arr.min()), 2),
                "储电量_E24_max": round(float(E24_arr.max()), 2),
            },
        },
        "fallback_trigger": {
            "fallback_id": "F4",
            "condition": "HD3/HD4口径变化或极端电价(0.0076)导致储能解退化",
            "observed": False,
            "evidence": None,
        },
    }
    with open(os.path.join(out_dir, "experiments", ROUND, "run_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


if __name__ == "__main__":
    main()
