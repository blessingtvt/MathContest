# -*- coding: utf-8 -*-
# Q1 主模型 M1: 代表日(附件1) 单日 LP, 最小化全天购电费 C1=Σ p_t·x_t, 终态 E_0=E_T=6000。
# 输出 result1.xlsx + experiments/round1/run_summary.json
import os, sys, json
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import numpy as np
import openpyxl
from common import *


def main():
    b = load_bundle()
    a1 = b["附件1"]
    price = a1["price"]; load = a1["load"]; pv = a1["pv_forecast"]

    # 主模型 M1
    cost, x, c, q, E = daily_lp(price, load, pv, E0=E0_INIT, terminal_eq=True, E_term=E0_INIT)
    # 基线
    b_cost, bx = no_storage(price, load, pv)
    r_cost, rx, rc, rq, rE = rule_arbitrage(price, load, pv, E0=E0_INIT)

    # ---- 写 result1.xlsx ----
    out_dir = os.path.join(RESULTS_DIR, "Q1")
    os.makedirs(os.path.join(out_dir, "experiments", "round1"), exist_ok=True)
    out_path = os.path.join(out_dir, "result1.xlsx")
    wb = openpyxl.load_workbook(os.path.join(TEMPLATE_DIR, "result1.xlsx"))
    ws = wb["计划购电量"]
    for t in range(T):
        ws.cell(row=2 + t, column=2, value=float(x[t]))
    ws2 = wb["充放电量"]
    for s in range(6):
        lo, hi = s * 24, (s + 1) * 24
        ws2.cell(row=2 + s, column=2, value=float(c[lo:hi].sum()))  # 充电量
        ws2.cell(row=2 + s, column=3, value=float(q[lo:hi].sum()))  # 放电量
    ws2.cell(row=2, column=5, value=float(E[0]))    # 0:00 储电量
    ws2.cell(row=3, column=5, value=float(E[-1]))   # 24:00 储电量
    wb.save(out_path)

    # ---- run_summary.json ----
    summary = {
        "schema_version": 1,
        "question_id": "Q1",
        "experiment_id": "round1",
        "method": "M1-LP",
        "seed": SEED,
        "status": "success",
        "created_at": datetime.now().isoformat(),
        "data_sources": ["data_clean/cleaned_data.pkl"],
        "terminal": {"E0": E0_INIT, "E_T": float(E[-1]), "constraint": "E0=ET"},
        "main": {
            "method_id": "M1",
            "objective_name": "全天购电费(元)",
            "objective_value": round(cost, 4),
            "total_purchase_kwh": round(float(x.sum()), 4),
            "total_charge_kwh": round(float(c.sum()), 4),
            "total_discharge_kwh": round(float(q.sum()), 4),
            "outputs": ["results/Q1/result1.xlsx"],
        },
        "baselines": [
            {"id": "B1", "name": "无储能", "objective_value": round(b_cost, 4)},
            {"id": "R1", "name": "规则套利", "objective_value": round(r_cost, 4),
             "terminal_E_kwh": round(float(rE[-1]), 4)},
        ],
        "metrics": {
            "储能价值_元": round(b_cost - cost, 4),
            "储能价值_占比": round((b_cost - cost) / b_cost, 6),
            "优化价值_元": round(r_cost - cost, 4),
        },
    }
    with open(os.path.join(out_dir, "experiments", "round1", "run_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


if __name__ == "__main__":
    main()
